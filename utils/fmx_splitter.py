"""Extract and number individual X-ray images from composite FMX sheets.

Handles two input formats:
  1. Composite image (single image with multiple X-rays on a dark background)
  2. PPTX file with individually placed images or composite slides

Each detected X-ray is numbered 1–N, row-by-row left-to-right, top-to-bottom,
matching the standard FMX layout order.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import List, Union

import cv2
import numpy as np
from PIL import Image


def _cluster_rows(boxes: List[dict], expected_rows: int = 0) -> List[List[dict]]:
    """Group bounding boxes into rows by finding the largest y-gaps.

    Uses the y-top coordinate for each box. Sorts all boxes by y-top, computes
    gaps between consecutive values, and splits at the largest gaps.

    If expected_rows is 0, auto-detects by looking for gaps significantly larger
    than the median gap.
    """
    if not boxes:
        return []

    sorted_boxes = sorted(boxes, key=lambda b: b["y"])

    if len(sorted_boxes) <= 1:
        return [sorted_boxes]

    # Compute gaps between consecutive y-tops
    gaps = []
    for i in range(1, len(sorted_boxes)):
        gaps.append((sorted_boxes[i]["y"] - sorted_boxes[i - 1]["y"], i))

    gaps.sort(reverse=True)

    if expected_rows > 1:
        num_splits = expected_rows - 1
    else:
        # Auto-detect: find gaps that are at least 2x the median gap
        gap_values = sorted([g for g, _ in gaps], reverse=True)
        if len(gap_values) >= 2:
            median_gap = sorted(gap_values)[len(gap_values) // 2]
            large_gap_threshold = max(median_gap * 2, gap_values[0] * 0.4)
            num_splits = sum(1 for g in gap_values if g >= large_gap_threshold)
            num_splits = max(1, min(num_splits, len(sorted_boxes) - 1))
        else:
            num_splits = 1

    # Get the split indices (the largest gaps)
    split_indices = sorted([idx for _, idx in gaps[:num_splits]])

    # Build rows
    rows = []
    prev = 0
    for si in split_indices:
        rows.append(sorted_boxes[prev:si])
        prev = si
    rows.append(sorted_boxes[prev:])

    # Sort each row left-to-right by x
    for row in rows:
        row.sort(key=lambda b: b["x"])

    return rows


def split_composite_image(
    image: Union[Image.Image, np.ndarray],
    threshold: int = 25,
    min_area_fraction: float = 0.003,
    padding: int = 2,
) -> List[dict]:
    """Detect individual X-ray regions in a composite FMX image.

    Parameters
    ----------
    image : PIL Image or numpy array (RGB/grayscale)
    threshold : brightness threshold to separate X-rays from background
    min_area_fraction : ignore regions smaller than this fraction of total area
    padding : pixels to add around each detected region when cropping

    Returns
    -------
    List of dicts sorted by position, each containing:
        number : int (1-based)
        image  : PIL.Image.Image (cropped X-ray)
        bbox   : tuple (x, y, w, h) in original image coordinates
        row    : int (0-based row index)
        col    : int (0-based column index within the row)
    """
    if isinstance(image, Image.Image):
        img_np = np.array(image.convert("RGB"))
    else:
        img_np = image

    gray = cv2.cvtColor(img_np, cv2.COLOR_RGB2GRAY) if img_np.ndim == 3 else img_np

    # Threshold to find bright X-ray regions on dark background
    _, thresh_img = cv2.threshold(gray, threshold, 255, cv2.THRESH_BINARY)

    # Clean up with morphological operations
    kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (5, 5))
    thresh_img = cv2.morphologyEx(thresh_img, cv2.MORPH_CLOSE, kernel, iterations=3)
    thresh_img = cv2.morphologyEx(thresh_img, cv2.MORPH_OPEN, kernel, iterations=2)

    # Find contours
    contours, _ = cv2.findContours(thresh_img, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

    img_area = gray.shape[0] * gray.shape[1]
    min_area = img_area * min_area_fraction

    boxes = []
    for c in contours:
        if cv2.contourArea(c) < min_area:
            continue
        x, y, w, h = cv2.boundingRect(c)
        boxes.append({"x": x, "y": y, "w": w, "h": h})

    if not boxes:
        return []

    # Cluster into rows, then number left-to-right, top-to-bottom
    rows = _cluster_rows(boxes)

    results = []
    number = 1
    h_img, w_img = gray.shape[:2]
    for row_idx, row in enumerate(rows):
        for col_idx, box in enumerate(row):
            x, y, w, h = box["x"], box["y"], box["w"], box["h"]
            # Apply padding (clamped to image bounds)
            x0 = max(0, x - padding)
            y0 = max(0, y - padding)
            x1 = min(w_img, x + w + padding)
            y1 = min(h_img, y + h + padding)
            crop = img_np[y0:y1, x0:x1]
            results.append({
                "number": number,
                "image": Image.fromarray(crop),
                "bbox": (x, y, w, h),
                "row": row_idx,
                "col": col_idx,
            })
            number += 1

    return results


def extract_from_pptx(pptx_path: Union[str, Path]) -> List[dict]:
    """Extract FMX images from a PowerPoint file.

    Handles two slide formats:
      - Slides with a single large composite image → runs split_composite_image
      - Slides with many individually-placed picture shapes → sorts by position

    Returns
    -------
    List of dicts per slide, each containing:
        slide_number : int
        title        : str (text from the slide, if any)
        images       : list of dicts from split_composite_image or shape extraction
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(pptx_path))
    slides_data = []

    for slide_idx, slide in enumerate(prs.slides):
        # Collect pictures and text
        pictures = []
        title = ""
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures.append(shape)
            elif hasattr(shape, "text") and shape.text.strip():
                if not title:
                    title = shape.text.strip()

        if not pictures:
            continue

        # Decide: single composite vs. multiple individual images
        if len(pictures) == 1:
            # Single image — treat as composite, split it
            blob = pictures[0].image.blob
            img = Image.open(io.BytesIO(blob)).convert("RGB")
            numbered = split_composite_image(img)
        elif len(pictures) >= 10:
            # Many images — individually placed, sort by position on slide
            boxes = []
            for pic in pictures:
                blob = pic.image.blob
                img = Image.open(io.BytesIO(blob)).convert("RGB")
                boxes.append({
                    "x": pic.left,
                    "y": pic.top,
                    "w": pic.width,
                    "h": pic.height,
                    "pil_image": img,
                })

            rows = _cluster_rows(boxes)
            numbered = []
            number = 1
            for row_idx, row in enumerate(rows):
                for col_idx, box in enumerate(row):
                    numbered.append({
                        "number": number,
                        "image": box["pil_image"],
                        "bbox": (box["x"], box["y"], box["w"], box["h"]),
                        "row": row_idx,
                        "col": col_idx,
                    })
                    number += 1
        else:
            # Few images but more than 1 — treat each as a separate composite
            numbered = []
            for pic in pictures:
                blob = pic.image.blob
                img = Image.open(io.BytesIO(blob)).convert("RGB")
                sub = split_composite_image(img)
                if sub:
                    # Re-number from current count
                    offset = len(numbered)
                    for s in sub:
                        s["number"] += offset
                    numbered.extend(sub)
                else:
                    numbered.append({
                        "number": len(numbered) + 1,
                        "image": img,
                        "bbox": (pic.left, pic.top, pic.width, pic.height),
                        "row": 0,
                        "col": len(numbered),
                    })

        slides_data.append({
            "slide_number": slide_idx + 1,
            "title": title,
            "images": numbered,
        })

    return slides_data


def save_extracted_images(
    images: List[dict],
    output_dir: Union[str, Path],
    prefix: str = "xray",
) -> List[Path]:
    """Save numbered X-ray images to a directory.

    Files are named {prefix}_{number:02d}.png (e.g. xray_01.png).
    Returns list of saved file paths.
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    saved = []
    for item in images:
        fname = f"{prefix}_{item['number']:02d}.png"
        fpath = output_dir / fname
        item["image"].save(str(fpath))
        saved.append(fpath)

    return saved
